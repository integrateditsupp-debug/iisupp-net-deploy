from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs" / "selected-pursuits-2026-06-05"
UNCOMMON_SRC = ROOT / "procurement-downloads" / "uncommon-2026-06-usi" / "UncommonSchools_RFP-2026-06-USI_CostProposal.xlsx"


HEADER = PatternFill("solid", fgColor="1F4D78")
SUBHEADER = PatternFill("solid", fgColor="D9EAF7")
WARN = PatternFill("solid", fgColor="FFF2CC")
GOOD = PatternFill("solid", fgColor="D9EAD3")
BAD = PatternFill("solid", fgColor="F4CCCC")
WHITE = "FFFFFF"
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_sheet(ws):
    ws.freeze_panes = "A2"
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
    for cell in ws[1]:
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = HEADER
    widths = {
        "A": 28,
        "B": 26,
        "C": 18,
        "D": 18,
        "E": 18,
        "F": 18,
        "G": 34,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def money(n):
    return n


def build_pricing_profit_model():
    wb = Workbook()
    ws = wb.active
    ws.title = "Pricing and Profit"
    rows = [
        [
            "Opportunity",
            "Pricing Ask",
            "Estimated Direct Cost",
            "Gross Profit",
            "Gross Margin",
            "Submission Status",
            "Notes",
        ],
        [
            "Mahone Bay base website redesign",
            money(68500),
            money(39000),
            "=B2-C2",
            "=D2/B2",
            "Ready after refs/signature/insurance",
            "Fixed project price excluding HST. Direct cost assumes CMS/accessibility/content bench plus IIS PM/QA.",
        ],
        [
            "Mahone Bay standard support - annual",
            money(14400),
            money(7200),
            "=B3-C3",
            "=D3/B3",
            "Optional recurring",
            "CAD 1,200/month excluding HST.",
        ],
        [
            "Mahone Bay enhanced support - annual",
            money(22200),
            money(10800),
            "=B4-C4",
            "=D4/B4",
            "Optional recurring",
            "CAD 1,850/month excluding HST.",
        ],
        [
            "Mahone Bay AI content/search advisory",
            money(4800),
            money(1600),
            "=B5-C5",
            "=D5/B5",
            "Optional add-on",
            "Fixed add-on excluding HST.",
        ],
        [
            "CanadaBuys AI QuickStart",
            money(18500),
            money(7400),
            "=B6-C6",
            "=D6/B6",
            "Catalog/rate-card item",
            "USD from Sourcewell rate card; convert/quote CAD as required by future callups.",
        ],
        [
            "CanadaBuys AI Pilot Sprint - Standard",
            money(58500),
            money(26325),
            "=B7-C7",
            "=D7/B7",
            "Catalog/rate-card item",
            "USD from Sourcewell rate card.",
        ],
        [
            "CanadaBuys Workflow Automation Build",
            money(42500),
            money(17000),
            "=B8-C8",
            "=D8/B8",
            "Catalog/rate-card item",
            "USD from Sourcewell rate card.",
        ],
        [
            "CanadaBuys Managed AI Ops - Essential annual",
            money(114000),
            money(45600),
            "=B9-C9",
            "=D9/B9",
            "Catalog/rate-card item",
            "USD 9,500/month from Sourcewell rate card.",
        ],
        [
            "Uncommon IIS workstream - transition",
            money(85000),
            money(34000),
            "=B10-C10",
            "=D10/B10",
            "Partner/subcontract target",
            "IIS lane only: AI service desk automation, M365/identity documentation, transition playbooks.",
        ],
        [
            "Uncommon IIS workstream - annual recurring",
            money(456000),
            money(228000),
            "=B11-C11",
            "=D11/B11",
            "Partner/subcontract target",
            "CAD/USD to be negotiated with prime. Assumes 38,000/month IIS workstream.",
        ],
        [
            "Uncommon full prime draft - Year 1",
            money(7888000),
            money(6626000),
            "=B12-C12",
            "=D12/B12",
            "Do not submit solo",
            "Only valid with named compliant prime/partners and proof. High revenue but operational/legal risk if IIS submits alone.",
        ],
    ]
    for row in rows:
        ws.append(row)
    style_sheet(ws)
    for row in range(2, ws.max_row + 1):
        ws[f"B{row}"].number_format = '$#,##0'
        ws[f"C{row}"].number_format = '$#,##0'
        ws[f"D{row}"].number_format = '$#,##0'
        ws[f"E{row}"].number_format = '0.0%'
        status = ws[f"F{row}"].value or ""
        if "Ready" in status or "Optional" in status or "Catalog" in status:
            ws[f"F{row}"].fill = GOOD
        if "Do not" in status:
            ws[f"F{row}"].fill = BAD
        if "Partner" in status:
            ws[f"F{row}"].fill = WARN

    ws2 = wb.create_sheet("Rate Card")
    ws2.append(["Role / Package", "NTE Rate", "Unit", "Estimated Cost Basis", "Expected Margin", "Notes"])
    rate_rows = [
        ["Executive AI Strategy Lead", 275, "USD/hour", 110, "60%", "Sourcewell rate card."],
        ["AI Solution Architect", 240, "USD/hour", 115, "52%", "Sourcewell rate card."],
        ["Responsible AI / Governance Lead", 225, "USD/hour", 100, "56%", "Sourcewell rate card."],
        ["AI Automation Engineer", 195, "USD/hour", 95, "51%", "Sourcewell rate card."],
        ["Microsoft 365 / Azure AI Specialist", 190, "USD/hour", 90, "53%", "Sourcewell rate card."],
        ["ITSM / Support Automation Specialist", 175, "USD/hour", 80, "54%", "Sourcewell rate card."],
        ["Business Analyst / Process Analyst", 150, "USD/hour", 70, "53%", "Sourcewell rate card."],
        ["Trainer / Change Lead", 145, "USD/hour", 65, "55%", "Sourcewell rate card."],
        ["Project Manager", 165, "USD/hour", 80, "52%", "Sourcewell rate card."],
        ["Technical Writer", 125, "USD/hour", 55, "56%", "Sourcewell rate card."],
    ]
    for row in rate_rows:
        ws2.append(row)
    style_sheet(ws2)
    ws2.column_dimensions["A"].width = 36

    path = OUT / "IIS_Selected_Pursuits_Pricing_and_Profit_Model.xlsx"
    wb.save(path)
    return path


def build_submission_control():
    path = OUT / "FINAL_SUBMISSION_CONTROL.md"
    path.write_text(
        """# Final Submission Control

Created: 2026-06-05

## Executive Decision

Proceed with all three pursuits, but submit only what is truthful and supportable.

- Mahone Bay: finalizable as a direct bid after references, insurance path, and Ahmad signature.
- CanadaBuys AI Source List: finalizable after SAP Business Network forms and mandatory criteria are downloaded and mapped.
- Uncommon Schools: finalizable only as partner-backed prime/team bid or subcontract/workstream package. Do not submit standalone prime without proof.

## Prices We Are Asking

### Mahone Bay

- Base website redesign: CAD 68,500 excluding HST.
- Standard support: CAD 1,200/month, CAD 14,400/year excluding HST.
- Enhanced support: CAD 1,850/month, CAD 22,200/year excluding HST.
- Optional AI content/search advisory: CAD 4,800 excluding HST.
- Additional services: CAD 125-CAD 175/hour excluding HST.

Expected gross profit:

- Base redesign: about CAD 29,500 gross profit, about 43% margin if direct/subcontract costs stay near CAD 39,000.
- Standard support: about CAD 7,200/year gross profit, about 50% margin.
- Enhanced support: about CAD 11,400/year gross profit, about 51% margin.
- Optional AI advisory: about CAD 3,200 gross profit, about 67% margin.

### CanadaBuys AI Source List

This is a qualification/source-list vehicle, not a single priced project. Use the Sourcewell AI rate card as the starting catalog:

- AI Readiness QuickStart: USD 18,500.
- Responsible AI Governance Toolkit: USD 24,500.
- AI Pilot Sprint - Standard: USD 58,500.
- AI Pilot Sprint - Advanced: USD 98,000.
- Workflow Automation Build: USD 42,500.
- Managed AI Operations - Essential: USD 9,500/month.
- Managed AI Operations - Standard: USD 18,500/month.
- Managed AI Operations - Enterprise: USD 37,500/month.
- AI Training Workshop: USD 6,500/day.

Expected gross profit:

- Target 45%-60% gross margin on advisory, automation, training, and managed AI operations if delivered by IIS plus controlled subcontract bench.

### Uncommon Schools

Recommended IIS target: partner/subcontract workstream, not standalone prime.

- IIS transition/workstream setup: USD/CAD 85,000 target.
- IIS recurring AI/M365/ITSM automation workstream: USD/CAD 38,000/month, USD/CAD 456,000/year target.

Expected gross profit:

- Transition workstream: about 60% gross margin if delivered by IIS and a small support bench.
- Recurring workstream: about 50% gross margin before partner/prime markups.

Full prime draft economics, not recommended solo:

- Year 1 full-service draft price: about USD/CAD 7.888M.
- Estimated direct cost: about USD/CAD 6.626M.
- Estimated gross profit: about USD/CAD 1.262M, about 16% margin.
- This is not submit-ready without a compliant MSP/onsite/SOC/state-licensed partner.

## Submission Options

### Mahone Bay

Portal: https://townofmahonebay.bidsandtenders.ca/Module/Tenders/en/Tender/Detail/d34bda17-ea5d-407f-997f-5b269a9b1e78

Upload:

- `IIS_MahoneBay_RFP2_Website_Redesign_Proposal.pdf`
- Signed Appendix A if the portal requires a separate signed form.
- Any reference appendix added before final upload.

Stop if the portal asks for legal/insurance confirmations Ahmad has not verified.

### CanadaBuys AI Source List

Portal: https://canadabuys.canada.ca/en/tender-opportunities/tender-notice/ws4286933967-doc4822970058

Go through SAP Business Network from the notice.

Upload/use:

- `IIS_CanadaBuys_AI_Source_List_Qualification_Package.pdf`
- Sourcewell pricing/rate-card material if the SAP event requests service catalog/rates.
- Portal copy from `PORTAL_PASTE_BLOCKS.md`.

Stop if the event asks for unsupported security certifications, insurance, references, or legal declarations.

### Uncommon Schools

Portal/page: https://uncommonschools.org/current-legal-notices/

Submission email in RFP: msp.rfp.2026@uncommonschools.org

Use only after partner/proof gate:

- `IIS_Uncommon_RFP2026_06_USI_Pursuit_and_Partner_Package.pdf`
- Draft cost workbook only if partner-backed and final-approved.
- New Jersey required forms only after Ahmad signs/attests.

Do not submit standalone prime unless the mandatory proof gaps are resolved.
""",
        encoding="utf-8",
    )
    return path


def build_uncommon_draft_cost_workbook():
    wb = load_workbook(UNCOMMON_SRC)
    # Conservative high-level full prime draft. This is intentionally marked as partner-gated.
    y1 = {
        "A - Enterprise Managed": {
            "B5": 32000,
            "B6": 24000,
            "B7": 12000,
            "B9": 45000,
            "B10": 28000,
            "B11": 18000,
            "B12": 22000,
            "B14": 28000,
            "B15": 24000,
            "B16": 20000,
            "B18": 225000,
        },
        "B - End User Support": {
            "B4": 85000,
            "B5": 160000,
            "B6": 48000,
            "B7": 38000,
            "B8": 22000,
            "B10": 175000,
        },
        "Project Mgmt": {
            "B6": 18000,
            "B8": 185,
            "B9": 245,
            "B10": 285,
        },
    }
    for sheet, cells in y1.items():
        ws = wb[sheet]
        for addr, val in cells.items():
            ws[addr] = val
        # 3% annual escalation
        for row in ws.iter_rows():
            if row[1].value and isinstance(row[1].value, (int, float)):
                row[2].value = round(row[1].value * 1.03, 2)
                row[3].value = round(row[1].value * 1.03**2, 2)
                row[4].value = round(row[1].value * 1.03**3, 2)
                row[5].value = round(row[1].value * 1.03**4, 2)
    off = wb["Offboarding"]
    for r in [5, 6, 7]:
        off[f"B{r}"] = "Y"
        off[f"C{r}"] = 0
        off[f"D{r}"] = "Included in base recurring fees, subject to final partner/legal approval."

    multi = wb["Multi-Year Pricing"]
    for row in [6, 7, 8]:
        for col in range(2, 6):
            multi.cell(row=row, column=col).value = 0.03
            multi.cell(row=row, column=col).number_format = "0.00%"
    multi["B10"] = "3% annual cap used for draft pricing; final escalation subject to partner approval."

    ass = wb["Assumptions"]
    assumptions = {
        5: "Draft partner-backed staffing model: IIS owns AI-enabled service desk automation, M365/identity workflow, documentation, reporting, transition playbooks, and QA. Regional onsite field technicians, NOC/SOC, school coverage, and state-compliant resources must be supplied by named partners before final submission.",
        6: "Draft assumes 24x7 monitoring/NOC through partner, help desk coverage aligned to RFP requirements, and IIS escalation/documentation support during agreed business and project windows.",
        7: "Draft assumes all 63 locations are covered through regional partner onsite teams in NY, NJ, and MA. Onsite frequency must be finalized by partner and Uncommon service baseline.",
        8: "Pricing is predicated on baseline stated in RFP/addenda: approximately 2,800 staff, 20,000 students, 21,800 devices, and 63 locations.",
        9: "Includes enterprise managed services, end user support, transition documentation, M365/identity support, AI-enabled knowledge/runbook automation, service reporting, and project/change support as scoped in final partner response.",
        10: "Excludes unsupported legal attestations, unverified state licensing, unsupported SOC/ISO claims, pass-through hardware/software/cloud/license costs, and any work requiring proof not available at proposal time.",
        11: "Cybersecurity tooling must be supplied/confirmed by SOC/MDR partner. IIS will not claim SOC 2/ISO coverage unless partner proof is attached.",
        12: "Draft uses 3% annual escalation for Years 2-5. Final escalation to be confirmed by prime/partner and legal review.",
        13: "Offboarding obligations are included in base recurring fees in the draft; final contract language must be reviewed before signature.",
        14: "Key subcontractors are required before final submission: regional education MSP/onsite partner, SOC/MDR provider, and any state-compliance resources.",
        17: "This workbook is a draft pricing model only and is not submit-ready until partner, insurance, licensing, certification, reference, and legal proof gates are resolved.",
    }
    for row, text in assumptions.items():
        wb["Assumptions"][f"C{row}"] = text

    path = OUT / "DRAFT_DO_NOT_SUBMIT_UncommonSchools_CostProposal_PartnerBacked.xlsx"
    wb.save(path)
    return path


def main():
    for p in [
        build_pricing_profit_model(),
        build_submission_control(),
        build_uncommon_draft_cost_workbook(),
    ]:
        print(p)


if __name__ == "__main__":
    main()
