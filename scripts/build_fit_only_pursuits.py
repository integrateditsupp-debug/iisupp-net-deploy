from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs" / "fit-only-pursuits-2026-06-05"
OUT.mkdir(parents=True, exist_ok=True)

COMPANY = {
    "legal_name": "Integrated IT Support Inc.",
    "contact": "Ahmad Wasee",
    "title": "Founder and Senior AI Engineer",
    "email": "ahmad.wasee@iisupp.net",
    "phone": "647-581-3182",
    "address": "30 Fothergill Court, Whitby, Ontario L1P 1L4, Canada",
    "website": "https://iisupp.net",
}


def style_doc(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for name, size, color in [
        ("Heading 1", 16, RGBColor(46, 116, 181)),
        ("Heading 2", 13, RGBColor(46, 116, 181)),
        ("Heading 3", 12, RGBColor(31, 77, 120)),
    ]:
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = color
        st.paragraph_format.space_before = Pt(10)
        st.paragraph_format.space_after = Pt(5)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(str(text))
    r.font.name = "Calibri"
    r.font.size = Pt(9.5)
    r.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def title(doc, text, subtitle):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor(31, 77, 120)
    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = s.add_run(subtitle)
    sr.italic = True
    sr.font.size = Pt(11)


def kv(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_text(table.rows[0].cells[0], "Field", True)
    cell_text(table.rows[0].cells[1], "Response", True)
    shade(table.rows[0].cells[0], "F2F4F7")
    shade(table.rows[0].cells[1], "F2F4F7")
    for k, v in rows:
        c = table.add_row().cells
        cell_text(c[0], k, True)
        cell_text(c[1], v)
    doc.add_paragraph()


def bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def rate_table(doc):
    doc.add_heading("Proposed Rate Card", level=1)
    rows = [
        ("Executive AI / Digital Strategy Lead", "CAD 245/hour", "Senior advisory, roadmap, governance, stakeholder workshops."),
        ("AI Automation Engineer", "CAD 185/hour", "Workflow automation, prompt/RAG/agent design, implementation support."),
        ("M365 / Azure / SharePoint Specialist", "CAD 180/hour", "M365, Entra, Intune, SharePoint, Teams, cloud workflow support."),
        ("Data / Knowledge Management Consultant", "CAD 175/hour", "Analytics, reporting, knowledge architecture, data governance support."),
        ("ITSM / Service Desk Process Specialist", "CAD 165/hour", "ServiceNow-style process, SLAs/KPIs, documentation, support automation."),
        ("Business Analyst / Process Analyst", "CAD 145/hour", "Requirements, process maps, operating model, UAT support."),
        ("Project Manager", "CAD 155/hour", "Delivery coordination, risk tracking, reporting, schedule management."),
        ("Technical Writer / Trainer", "CAD 120/hour", "SOPs, user guides, training materials, knowledge base articles."),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for idx, h in enumerate(["Role", "Rate", "Use"]):
        cell_text(table.rows[0].cells[idx], h, True)
        shade(table.rows[0].cells[idx], "F2F4F7")
    for row in rows:
        c = table.add_row().cells
        for i, v in enumerate(row):
            cell_text(c[i], v, i == 0)
    doc.add_paragraph("Rates exclude applicable taxes and approved travel/pass-through costs unless a future SOW states otherwise.")


def build_sask():
    doc = Document()
    style_doc(doc)
    title(doc, "Qualification Response Shell", "SaskBuilds RFSQ_0586-1 - Information Management Technology Consulting Services")
    kv(
        doc,
        [
            ("Supplier", COMPANY["legal_name"]),
            ("Primary consultant", f"{COMPANY['contact']}, {COMPANY['title']}"),
            ("Contact", f"{COMPANY['email']} | {COMPANY['phone']}"),
            ("Address", COMPANY["address"]),
            ("Website", COMPANY["website"]),
            ("Opportunity status", "Open continuous prequalification; closes 2028-02-29 2:00 PM CST."),
            ("Submission address shown", "response@gov.sk.ca"),
        ],
    )
    doc.add_heading("Fit Summary", level=1)
    doc.add_paragraph(
        "This is a clean-fit pursuit for IIS because the public notice seeks prequalified vendors for information management technology consulting services, especially operational reporting, analytics, data management, and advisory support on an as-required basis."
    )
    bullets(
        doc,
        [
            "Ahmad has 13+ years of enterprise IT, ITSM, reporting, SLA/KPI, CMDB, documentation, M365, and regulated support experience.",
            "IIS operates ARIA, a practical AI and automation platform for knowledge retrieval, procurement intelligence, workflow automation, and operational briefings.",
            "The opportunity is a qualification list, not a heavy fixed-delivery bid, which matches current IIS capacity.",
            "No public synopsis requirement found for defence clearance, SOC/ISO at submission, multi-site onsite staffing, reseller authorization, or municipal website references.",
        ],
    )
    doc.add_heading("IIS Service Categories to Offer", level=1)
    bullets(
        doc,
        [
            "Operational reporting and analytics advisory.",
            "Data management, knowledge architecture, and AI-ready documentation.",
            "ServiceNow-style ITSM process review, KPI/SLA dashboard support, and support workflow optimization.",
            "Microsoft 365, SharePoint, Teams, Entra, Intune, and Azure-aligned workflow support.",
            "AI-assisted business process automation, document automation, and analyst enablement.",
            "Training, documentation, UAT, change support, and project coordination.",
        ],
    )
    doc.add_heading("Evidence Mapping", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for i, h in enumerate(["Requirement Theme", "IIS Evidence", "Proof Source"]):
        cell_text(table.rows[0].cells[i], h, True)
        shade(table.rows[0].cells[i], "F2F4F7")
    rows = [
        ("Operational reporting / analytics", "RBC SLA/KPI dashboards, reporting automation, ServiceNow/HP Service Manager queue analytics.", "Ahmad resume."),
        ("Data management / quality", "CMDB lifecycle management, CSDM alignment, evidence-of-control posture, knowledge-base architecture.", "Raymond James/RBC experience."),
        ("Information management", "SharePoint/M365, documentation, runbooks, KBs, onboarding guides, support workflow design.", "Enterprise IT background."),
        ("AI and automation", "ARIA platform, multi-agent automation, prompt/RAG workflows, AI-assisted triage and incident summaries.", "IIS/ARIA and Raymond James."),
        ("Government/public context", "Ontario Health, City of Toronto, regulated financial services, healthcare privacy practices.", "Ahmad resume."),
    ]
    for row in rows:
        c = table.add_row().cells
        for i, v in enumerate(row):
            cell_text(c[i], v, i == 0)
    doc.add_paragraph()
    rate_table(doc)
    doc.add_heading("Finalization Checklist", level=1)
    bullets(
        doc,
        [
            "Download the full RFSQ document and Appendix I rated criteria from SaskTenders account.",
            "Map mandatory/rated criteria to this response shell.",
            "Confirm whether references are required and use only permission-safe contacts.",
            "Confirm legal/insurance declarations before signature.",
            "Submit through the official response method only after Ahmad review.",
        ],
    )
    path = OUT / "IIS_SaskBuilds_RFSQ0586_IM_IT_Consulting_Response_Shell.docx"
    doc.save(path)
    return path


def build_tic():
    doc = Document()
    style_doc(doc)
    title(doc, "Qualification Response Shell", "Transportation Investment Corporation - TIC TEC MUL RFQ 2025 42R IT Advisory Services")
    kv(
        doc,
        [
            ("Supplier", COMPANY["legal_name"]),
            ("Primary advisor", f"{COMPANY['contact']}, {COMPANY['title']}"),
            ("Contact", f"{COMPANY['email']} | {COMPANY['phone']}"),
            ("Address", COMPANY["address"]),
            ("Website", COMPANY["website"]),
            ("Opportunity status", "Open RFQ; closes 2027-07-30 5:00 PM EDT per MERX abstract."),
            ("Procurement contact shown", "procurement@ticorp.ca"),
        ],
    )
    doc.add_heading("Fit Summary", level=1)
    doc.add_paragraph(
        "This is a clean-fit pursuit for Ahmad/IIS based on the public notice: TI Corp is seeking individual IT advisors to support multiple projects, including EDRMS and SharePoint integration, on an if/as/when-required basis."
    )
    bullets(
        doc,
        [
            "Ahmad has direct SharePoint, M365, ITSM, documentation, ServiceNow, project support, and enterprise support leadership experience.",
            "The RFQ seeks individual advisors, which fits Ahmad as the named primary advisor without requiring a large onsite delivery bench.",
            "EDRMS/SharePoint integration, information management, workflow automation, and advisory support align strongly with IIS capabilities.",
            "No public abstract requirement found for SOC/ISO at proposal time, defence clearance, reseller authorization, or large onsite staffing.",
        ],
    )
    doc.add_heading("Advisor Services Offered", level=1)
    bullets(
        doc,
        [
            "SharePoint Online, Microsoft 365, Teams, Entra, Intune, and workflow advisory.",
            "EDRMS/information management advisory, content migration planning, taxonomy, permissions, and governance support.",
            "IT project advisory, requirements gathering, UAT, documentation, training, and change support.",
            "Service desk/ITSM process analysis, KPI/SLA reporting, operational readiness, and support transition planning.",
            "AI-assisted knowledge management, workflow automation, and document/runbook acceleration where approved.",
        ],
    )
    doc.add_heading("Experience Mapping", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for i, h in enumerate(["TI Corp Need", "IIS/Ahmad Fit", "Proof Source"]):
        cell_text(table.rows[0].cells[i], h, True)
        shade(table.rows[0].cells[i], "F2F4F7")
    rows = [
        ("SharePoint integration", "M365/SharePoint/Teams support and enterprise knowledge/documentation workflows.", "Resume and IIS materials."),
        ("EDRMS/information management", "Document management migrations/support, knowledge-base architecture, governance-aware support.", "Cavalluzzo, Raymond James, IIS/ARIA."),
        ("Individual IT advisory", "Senior technical analyst, AI engineer, project support, BA/process support, incident/process leadership.", "Raymond James/RBC/City of Toronto."),
        ("Project team augmentation", "Used to high-stakes regulated environments, cross-team collaboration, executive/trader-floor support, UAT/docs.", "Enterprise experience."),
        ("Automation and reporting", "VBA/VB.NET automation, SLA/KPI dashboards, AI-assisted workflows and ARIA.", "RBC/IIS."),
    ]
    for row in rows:
        c = table.add_row().cells
        for i, v in enumerate(row):
            cell_text(c[i], v, i == 0)
    doc.add_paragraph()
    rate_table(doc)
    doc.add_heading("Finalization Checklist", level=1)
    bullets(
        doc,
        [
            "Log into MERX and download the full RFQ document/response forms.",
            "Confirm mandatory forms, reference requirements, insurance wording, and any British Columbia business requirements.",
            "Submit Ahmad as the named individual advisor only where criteria match his actual resume.",
            "Do not add unapproved subcontractor or certification claims.",
        ],
    )
    path = OUT / "IIS_TIC_IT_Advisory_RFQ_Response_Shell.docx"
    doc.save(path)
    return path


def build_invest_ns():
    doc = Document()
    style_doc(doc)
    title(doc, "Standing Offer Response Shell", "Invest Nova Scotia - Professional and Consulting Services and Certain Goods")
    kv(
        doc,
        [
            ("Supplier", COMPANY["legal_name"]),
            ("Primary contact", f"{COMPANY['contact']}, {COMPANY['title']}"),
            ("Contact", f"{COMPANY['email']} | {COMPANY['phone']}"),
            ("Address", COMPANY["address"]),
            ("Website", COMPANY["website"]),
            ("Opportunity status", "Continuous onboarding; submissions due before 2027-02-15 5:00 PM AST."),
            ("Submission portal", "Invest Nova Scotia Euna Procurement / Bonfire portal."),
        ],
    )
    doc.add_heading("Fit Summary", level=1)
    doc.add_paragraph(
        "This standing offer is a clean-fit pursuit for IIS because it prequalifies suppliers for professional and consulting services to support short-term needs. IIS can offer targeted AI readiness, Microsoft 365 and workflow consulting, ITSM/process improvement, reporting, documentation, and project support without claiming capabilities outside the current evidence base."
    )
    bullets(
        doc,
        [
            "Ontario-based corporation with practical consulting and managed IT positioning.",
            "Senior advisor has 13+ years of enterprise IT, regulated support, ITSM, M365, reporting, documentation, and automation experience.",
            "IIS operates ARIA, a production AI/automation platform used for procurement intelligence, knowledge retrieval, workflow support, and operational briefings.",
            "Standing-offer format suits IIS because future statements of work can be scoped to exact capabilities rather than forcing a large one-time delivery bid.",
        ],
    )
    doc.add_heading("Service Categories to Offer", level=1)
    bullets(
        doc,
        [
            "AI readiness assessments, use-case discovery, governance templates, and practical adoption planning.",
            "Microsoft 365, SharePoint, Teams, Entra, Intune, and Copilot readiness advisory.",
            "Workflow automation, document automation, knowledge-base structure, and AI-assisted operational support.",
            "ITSM/service desk process improvement, SLA/KPI reporting, incident/problem/change workflow support.",
            "Data/reporting advisory, dashboard requirements, business analysis, UAT, documentation, and training.",
            "Project coordination and delivery support for short-term technology initiatives.",
        ],
    )
    doc.add_heading("Evidence Mapping", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for i, h in enumerate(["Invest NS Need", "IIS Fit", "Proof Source"]):
        cell_text(table.rows[0].cells[i], h, True)
        shade(table.rows[0].cells[i], "F2F4F7")
    rows = [
        ("Professional consulting", "Senior technical advisory, ITSM/process support, project delivery, executive/user support.", "Raymond James, RBC, City of Toronto, Ontario Health."),
        ("Technology sector support", "AI, M365, Azure, Entra, Intune, SharePoint, support automation, workflow improvement.", "Ahmad resume and IIS/ARIA."),
        ("Short-term support needs", "Scoped fixed-fee or hourly advisory packages, rapid assessments, documentation/training deliverables.", "IIS service model."),
        ("Export/business-growth alignment", "AI and automation support for businesses improving operations and customer support.", "IIS company positioning."),
        ("Responsible delivery", "No unsupported SOC/ISO/reseller/onsite claims; use scoped work and approved support bench only where needed.", "Fit-only rule."),
    ]
    for row in rows:
        c = table.add_row().cells
        for i, v in enumerate(row):
            cell_text(c[i], v, i == 0)
    doc.add_paragraph()
    rate_table(doc)
    doc.add_heading("Suggested Portal Profile Text", level=1)
    doc.add_paragraph(
        "Integrated IT Support Inc. is an Ontario-based applied AI, automation, and managed IT consulting firm. IIS helps organizations improve operations through practical AI readiness, Microsoft 365 and SharePoint advisory, workflow automation, ITSM and service desk process improvement, reporting, documentation, training, and project support. IIS is led by Ahmad Wasee, a senior AI engineer and enterprise IT professional with 13+ years of experience across regulated financial services, healthcare, public-sector support, legal, and enterprise service management environments."
    )
    doc.add_heading("Finalization Checklist", level=1)
    bullets(
        doc,
        [
            "Create or log into the Invest Nova Scotia Euna/Bonfire vendor account.",
            "Download the standing offer documents and identify exact service categories.",
            "Attach corporate profile, Ahmad resume, rate card, and capability statement as allowed.",
            "Confirm references, insurance, tax/legal declarations, and signature fields before submission.",
            "Submit only categories that match actual IIS/Ahmad evidence.",
        ],
    )
    path = OUT / "IIS_InvestNovaScotia_StandingOffer_Response_Shell.docx"
    doc.save(path)
    return path


def build_tic_invest_control():
    path = OUT / "TIC_AND_INVEST_NS_SUBMISSION_CONTROL.md"
    path.write_text(
        """# TIC and Invest Nova Scotia Submission Control

Created: 2026-06-05

## Status

Not submitted yet. Response shells are ready, but final submission needs official portal forms and Ahmad approval for any declarations.

## 2. TIC IT Advisory Services

Official page: https://www.merx.com/public/supplier/solicitations/notice/443257908088/abstract

Prepared files:

- `IIS_TIC_IT_Advisory_RFQ_Response_Shell.docx`
- `IIS_TIC_IT_Advisory_RFQ_Response_Shell.pdf`

What to do in MERX:

1. Log into MERX.
2. Open TIC TEC MUL RFQ 2025 42R IT Advisory Services.
3. Download full RFQ documents and response forms.
4. Confirm mandatory requirements, reference form, insurance/legal language, and submission method.
5. Use Ahmad as the named individual advisor for SharePoint, EDRMS, M365, ITSM, documentation, workflow automation, and AI-enabled knowledge work.
6. Submit only after Ahmad approves any legal declarations.

## 3. Invest Nova Scotia Standing Offer

Official page: https://investnovascotia.ca/invest-nova-scotia-standing-offer

Submission portal: https://investnovascotia.bonfirehub.ca

Prepared files:

- `IIS_InvestNovaScotia_StandingOffer_Response_Shell.docx`
- `IIS_InvestNovaScotia_StandingOffer_Response_Shell.pdf`

What to do in Euna/Bonfire:

1. Create/log into the Euna/Bonfire vendor account.
2. Open the Invest Nova Scotia Standing Offer submission.
3. Download the standing offer package.
4. Select only fit categories: AI readiness, IT consulting, M365/SharePoint, ITSM/process, reporting/analytics, workflow automation, documentation, training, and project support.
5. Upload the response shell/capability materials if allowed.
6. Submit only after Ahmad verifies references, insurance/legal declarations, and signature fields.
""",
        encoding="utf-8",
    )
    return path


def build_global_scan():
    path = OUT / "GLOBAL_FIT_ONLY_SCAN.md"
    path.write_text(
        """# Global Fit-Only Scan

Created: 2026-06-05

Filter used: show only opportunities that appear to fit IIS without hard blockers such as defence/security clearance, SOC/ISO at proposal time, heavy onsite staffing, reseller-only requirements, bid bonds, or mandatory references IIS does not have.

## Pursue Now

### Canada - SaskBuilds and Procurement RFSQ_0586-1

Source: https://sasktenders.ca/content/public/print.aspx?competitionId=75ed0837-1245-4774-bf3b-025f69be4130

Why it fits: information management technology consulting, operational reporting, analytics, data management, and prequalified consultant list. This aligns with Ahmad's ITSM/reporting/data/AI/M365 experience.

Next move: download full RFSQ package from SaskTenders account, complete rated criteria, and submit only truthful evidence.

### Canada - Transportation Investment Corporation TIC TEC MUL RFQ 2025 42R

Source: https://www.merx.com/public/supplier/solicitations/notice/443257908088/abstract

Why it fits: individual IT advisors for EDRMS, SharePoint integration, and IT advisory on an if/as/when-required basis. This aligns with Ahmad's M365/SharePoint, documentation, workflow, ITSM, and project support background.

Next move: log into MERX, download full forms, and qualify Ahmad as named individual advisor.

### Canada - Invest Nova Scotia Standing Offer INS2024-25-23

Source: https://investnovascotia.ca/invest-nova-scotia-standing-offer

Why it fits: continuous onboarding for professional and consulting services until 2027-02-15. Broad enough for AI readiness, digital workflow, M365, ITSM, reporting, documentation, and project support.

Next move: create Euna/Bonfire account and review service categories before deciding whether to apply.

### Canada - Destination BC Data Analytics / Digital Strategy RFQs

Source: https://www.destinationbc.ca/work-with-us/contractor-supplier/

Why it fits: active RFQ lists include Data Analytics and Management of Digital Strategy and Information Technology Projects. Good match for analytics, AI readiness, reporting, M365/SharePoint, documentation, and project management.

Next move: use BC Bid account to download RFQ documents and check rated criteria.

## Watch / Partner Only

### UK

Contracts Finder and Find a Tender are useful markets, but live searches found mostly closed/awarded AI/M365 items or frameworks that appear already live rather than open for direct supplier onboarding. Use UK as a watch/partner market until a below-threshold AI/M365/ITSM advisory opportunity appears without UK-local proof barriers.

Useful official portal: https://www.gov.uk/contracts-finder

### USA

Federal SAM.gov work usually requires SAM registration/UEI and often has set-aside/local compliance or security gates. State/local opportunities can fit, but current search results found closed/archive items or partner/reseller requirements. Use USA as watch/partner unless IIS obtains SAM/UEI and finds a no-clearance training/advisory RFQ.

### Dubai / UAE

AI and Microsoft 365 activity is strong, but public tenders often require UAE supplier/vendor registration or local program access. Dubai SME/GPP is a useful registration path, not a direct lead yet.

Useful portal: https://www.mysme.ae/

### Qatar

Qatar public-sector AI activity is strong, but current search returned strategic partnerships and enterprise programs rather than open fit-only supplier tenders. Treat Qatar as partner/outreach market for AI/M365 adoption, not a direct bid market yet.

## Do Not Show

- Defence, military, classified, or battlefield-related work.
- Anything requiring SOC 2/ISO 27001 proof unless IIS has a compliant partner named.
- Mandatory multi-state onsite staffing.
- Reseller/CSP-only bids.
- Website tenders requiring municipal references unless IIS has approved references or a partner.
""",
        encoding="utf-8",
    )
    return path


def main():
    for p in [build_sask(), build_tic(), build_invest_ns(), build_tic_invest_control(), build_global_scan()]:
        print(p)


if __name__ == "__main__":
    main()
