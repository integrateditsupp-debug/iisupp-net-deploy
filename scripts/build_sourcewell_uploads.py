from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    ListFlowable,
    ListItem,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs" / "sourcewell-rfp061726" / "upload-ready"
OUT.mkdir(parents=True, exist_ok=True)


NAVY = "1F4E79"
BLUE = "D9EAF7"
GRAY = "F3F6F8"
LINE = "B8C7D2"


def style_sheet():
    styles = getSampleStyleSheet()
    styles["Title"].fontName = "Helvetica-Bold"
    styles["Title"].fontSize = 18
    styles["Title"].leading = 22
    styles["Heading1"].fontName = "Helvetica-Bold"
    styles["Heading1"].fontSize = 13
    styles["Heading1"].leading = 16
    styles["Heading1"].spaceBefore = 12
    styles["Heading1"].spaceAfter = 6
    styles["BodyText"].fontName = "Helvetica"
    styles["BodyText"].fontSize = 9.5
    styles["BodyText"].leading = 13
    styles["BodyText"].spaceAfter = 6
    styles.add(
        ParagraphStyle(
            name="Small",
            parent=styles["BodyText"],
            fontSize=8.5,
            leading=11,
            alignment=TA_LEFT,
        )
    )
    return styles


def footer(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#4B5563"))
    canvas.drawString(0.75 * inch, 0.5 * inch, "Integrated IT Support Inc. | Sourcewell RFP 061726")
    canvas.drawRightString(7.75 * inch, 0.5 * inch, f"Page {doc.page}")
    canvas.restoreState()


def bullets(items, styles):
    return ListFlowable(
        [ListItem(Paragraph(item, styles["BodyText"]), leftIndent=12) for item in items],
        bulletType="bullet",
        start="circle",
        leftIndent=18,
    )


def build_pdf(path, title, sections):
    styles = style_sheet()
    doc = SimpleDocTemplate(
        str(path),
        pagesize=LETTER,
        rightMargin=0.75 * inch,
        leftMargin=0.75 * inch,
        topMargin=0.75 * inch,
        bottomMargin=0.75 * inch,
        title=title,
    )
    story = [Paragraph(title, styles["Title"]), Spacer(1, 0.12 * inch)]
    story.append(
        Paragraph(
            "Prepared for Sourcewell RFP 061726 - Enterprise Artificial Intelligence: Platforms, Services, and Integrated Delivery.",
            styles["BodyText"],
        )
    )
    for heading, content in sections:
        story.append(Paragraph(heading, styles["Heading1"]))
        if isinstance(content, list):
            story.append(bullets(content, styles))
        else:
            for para in content.split("\n\n"):
                story.append(Paragraph(para, styles["BodyText"]))
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


def build_pricing_workbook(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Pricing Summary"
    ws.append(["Integrated IT Support Inc."])
    ws.append(["Sourcewell RFP 061726 - Lot 2 AI Professional Services Pricing Schedule"])
    ws.append(["All prices are Sourcewell not-to-exceed USD prices. IIS may quote lower prices by scope, volume, timeline, and approved discounts."])
    ws.append([])
    ws.append(["Service Package", "Description", "Unit", "NTE USD", "Notes"])
    packages = [
        ["AI Readiness QuickStart", "2-week discovery, current-state review, use-case scoring, and 30/60/90 roadmap", "Fixed fee", 18500, "Entry advisory package"],
        ["Responsible AI Governance Toolkit", "Policy, acceptable use, risk intake, oversight workflow, and incident playbook", "Fixed fee", 24500, "Governance and controls"],
        ["AI Pilot Sprint - Standard", "4-6 week pilot for one workflow with UAT and launch plan", "Fixed fee", 58500, "Implementation package"],
        ["AI Pilot Sprint - Advanced", "8-10 week pilot with integrations, reporting, and change management", "Fixed fee", 98000, "Advanced implementation"],
        ["Workflow Automation Build", "Design/build for AI-enabled intake, routing, summarization, or knowledge retrieval", "Fixed fee", 42500, "Workflow automation"],
        ["Managed AI Operations - Essential", "Monitoring, prompt/workflow updates, reporting, and support queue review", "Monthly", 9500, "Recurring support"],
        ["Managed AI Operations - Standard", "Essential plus governance reviews, analytics, and improvement backlog", "Monthly", 18500, "Recurring support"],
        ["Managed AI Operations - Enterprise", "Standard plus expanded support, multiple workflows/models, and steering support", "Monthly", 37500, "Recurring support"],
        ["AI Training Workshop", "Half-day or full-day role-based AI training", "Day", 6500, "Training and adoption"],
    ]
    for row in packages:
        ws.append(row)

    rates = wb.create_sheet("Role Rate Card")
    rates.append(["Role", "NTE USD / Hour", "Notes"])
    for row in [
        ["Executive AI Strategy Lead", 275, "Senior workshops, roadmap, governance"],
        ["AI Solution Architect", 240, "Architecture, system design, model/vendor selection"],
        ["Responsible AI / Governance Lead", 225, "Policies, risk, controls, oversight"],
        ["AI Automation Engineer", 195, "Workflow implementation, prompt/RAG/agents"],
        ["Microsoft 365 / Azure AI Specialist", 190, "M365, Azure, Copilot, Entra, integration"],
        ["Data / Knowledge Engineer", 185, "Knowledge base, retrieval, data preparation"],
        ["ITSM / Support Automation Specialist", 175, "Service desk, ticketing, operations workflows"],
        ["Business Analyst / Process Analyst", 150, "Discovery, requirements, process design"],
        ["Trainer / Change Lead", 145, "Training, adoption, materials"],
        ["Project Manager", 165, "Delivery governance, reporting, schedule"],
        ["QA / UAT Analyst", 135, "Test plans, acceptance, quality checks"],
        ["Technical Writer", 125, "SOPs, user guides, admin docs"],
    ]:
        rates.append(row)

    discounts = wb.create_sheet("Discounts and Exclusions")
    discounts.append(["Annual Purchase Volume", "Proposed Discount"])
    for row in [
        ["Under USD 100,000", "0%"],
        ["USD 100,000 - 249,999", "3%"],
        ["USD 250,000 - 499,999", "5%"],
        ["USD 500,000 - 999,999", "8%"],
        ["USD 1,000,000+", "10% or negotiated project discount"],
    ]:
        discounts.append(row)
    discounts.append([])
    discounts.append(["Exclusions"])
    for item in [
        "Applicable taxes, duties, withholding, and government-imposed fees",
        "Third-party AI/model/API/platform licensing",
        "Cloud compute, storage, networking, model, or API consumption",
        "Hardware, travel, onsite expenses, specialized audits, and approved pass-through costs",
        "Customer environment licensing such as Microsoft, ServiceNow, CRM, analytics, or security tooling",
    ]:
        discounts.append([item])

    for sheet in wb.worksheets:
        sheet.freeze_panes = "A5" if sheet.title == "Pricing Summary" else "A2"
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.font = Font(name="Calibri", size=10)
        for cell in sheet[1]:
            cell.font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=NAVY)
        header_row = 5 if sheet.title == "Pricing Summary" else 1
        for cell in sheet[header_row]:
            cell.font = Font(name="Calibri", size=10, bold=True, color="1F2937")
            cell.fill = PatternFill("solid", fgColor=BLUE)
            cell.border = Border(bottom=Side(style="thin", color=LINE))
        for col in range(1, sheet.max_column + 1):
            letter = get_column_letter(col)
            width = 18
            if col == 1:
                width = 32
            elif col in (2, 5):
                width = 55
            sheet.column_dimensions[letter].width = width
        for row in range(1, sheet.max_row + 1):
            sheet.row_dimensions[row].height = 20
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '$#,##0'
    wb.save(path)


def main():
    build_pricing_workbook(OUT / "IIS_Sourcewell_RFP061726_Lot2_Pricing_Schedule.xlsx")

    build_pdf(
        OUT / "IIS_Sourcewell_RFP061726_Financial_Strength_and_Stability.pdf",
        "Financial Strength and Stability",
        [
            (
                "Company Profile",
                "Integrated IT Support Inc. (IIS) is a privately held, founder-led Canadian technology services company based in Whitby, Ontario. IIS is led by Ahmad Wasee, whose enterprise IT, support, automation, and AI experience includes financial services, healthcare/public-sector environments, municipal technology environments, managed support, and AI enablement work.",
            ),
            (
                "Financial Position",
                "IIS operates with a lean services model, low fixed overhead, and scope-controlled delivery. IIS does not rely on a large inventory model, long hardware supply chain, or capital-intensive platform ownership to deliver the Lot 2 services proposed. Work is structured through quotes, statements of work, milestone billing, monthly managed-service billing, and approved pass-through costs where applicable.",
            ),
            (
                "Delivery Resourcing",
                [
                    "IIS will remain the accountable supplier for Sourcewell participating entity work.",
                    "Ahmad Wasee will provide senior delivery leadership and technical oversight.",
                    "Approved specialist subcontractor or partner resources may be used for surge, niche security, cloud, training, accessibility, or integration needs under IIS governance.",
                    "Subcontractor use will be scoped, quality-controlled, and disclosed as required by the SOW and participating entity policies.",
                ],
            ),
            (
                "Risk Controls",
                [
                    "SOW-based scope, deliverables, assumptions, acceptance criteria, and change control.",
                    "No claim of bankruptcy, debarment, or government sales misrepresentation in the portal response.",
                    "Pricing excludes unapproved third-party software, cloud consumption, travel, taxes, and pass-through costs.",
                    "Client data ownership, confidentiality, and retention are handled through contract, SOW, and participating entity policy.",
                ],
            ),
        ],
    )

    build_pdf(
        OUT / "IIS_Sourcewell_RFP061726_Marketing_Plan_and_Samples.pdf",
        "Marketing Plan and Samples",
        [
            (
                "Sourcewell Marketing Approach",
                "If awarded, IIS will create a Sourcewell contract page on iisupp.net, publish a concise AI professional-services capability sheet, and promote the agreement through direct outreach, LinkedIn, webinar invitations, AI readiness offers, and public-sector-focused education content.",
            ),
            (
                "Sample Campaign Concepts",
                [
                    "No-cost 30-minute AI readiness triage for eligible Sourcewell participating entities.",
                    "Quarterly responsible AI and workflow automation webinars for municipalities, education, nonprofits, and public agencies.",
                    "One-page offer sheets for AI Readiness QuickStart, Responsible AI Governance Toolkit, AI Pilot Sprint, and Managed AI Operations.",
                    "Role-based AI literacy and secure-use training for agency staff and small organizations.",
                ],
            ),
            (
                "Sales Enablement",
                "IIS will train internal and approved partner resources on Sourcewell ordering, contract number use, quote/SOW templates, pricing rules, discount handling, administrative fee tracking, and participating entity eligibility. IIS will maintain CRM-style tracking for inquiries, proposals, orders, renewals, and required Sourcewell reporting.",
            ),
        ],
    )

    build_pdf(
        OUT / "IIS_Sourcewell_RFP061726_Standard_Transaction_Document_Samples.pdf",
        "Standard Transaction Document Samples",
        [
            (
                "Documents Used Under an Award",
                [
                    "Sourcewell contract quote referencing the master agreement number.",
                    "Statement of Work describing scope, deliverables, schedule, roles, assumptions, pricing, and acceptance criteria.",
                    "Project plan or implementation backlog for active delivery.",
                    "Support plan or service-level exhibit for managed services.",
                    "Change order for approved scope, timing, pricing, or responsibility changes.",
                    "Acceptance sign-off or delivery confirmation for milestone completion.",
                ],
            ),
            (
                "SOW Outline",
                "A typical IIS SOW includes: participating entity name; Sourcewell contract reference; project background; in-scope and out-of-scope work; deliverables; timeline; IIS and client responsibilities; data/security assumptions; pricing; invoice schedule; acceptance criteria; support model; change control; and required attachments.",
            ),
            (
                "Contract Hierarchy",
                "The Sourcewell master agreement will be treated as the controlling procurement framework. IIS transaction documents will be used to define operational details and will not override mandatory Sourcewell terms unless expressly permitted by Sourcewell and the participating entity.",
            ),
        ],
    )

    build_pdf(
        OUT / "IIS_Sourcewell_RFP061726_Responsible_AI_and_Governance_Support.pdf",
        "Responsible AI and Governance Support",
        [
            (
                "Governance Approach",
                "IIS uses a practical responsible AI delivery framework: use-case intake and scoring, data classification review, privacy/security review, acceptable-use rules, human oversight design, pilot approval gates, testing/UAT, documentation, handoff, and post-go-live monitoring.",
            ),
            (
                "Controls",
                [
                    "Participating entity data remains the participating entity's property.",
                    "No client data is used to train or fine-tune models unless expressly authorized in writing.",
                    "Sensitive data flows are reviewed before AI use.",
                    "Human-in-the-loop controls are recommended for consequential or regulated workflows.",
                    "Known limitations, escalation paths, and output verification responsibilities are documented.",
                    "Exit rights and data portability are addressed in the SOW and aligned to the master agreement.",
                ],
            ),
            (
                "Resourced Delivery",
                "IIS will provide senior-led AI professional services and may use approved specialist resources for cloud, security, accessibility, training, integration, or surge needs. IIS remains accountable for scope, quality, communication, and delivery controls.",
            ),
        ],
    )


if __name__ == "__main__":
    main()
