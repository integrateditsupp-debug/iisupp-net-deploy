from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs" / "selected-pursuits-2026-06-05"
OUT.mkdir(parents=True, exist_ok=True)

COMPANY = {
    "legal_name": "Integrated IT Support Inc.",
    "brand": "Integrated IT Support Inc. (IIS)",
    "address": "30 Fothergill Court, Whitby, Ontario L1P 1L4, Canada",
    "contact": "Ahmad Wasee",
    "title": "Founder and Senior AI Engineer",
    "email": "ahmad.wasee@iisupp.net",
    "phone": "647-581-3182",
    "website": "https://iisupp.net",
}


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = "Calibri"
    r.font.size = Pt(9.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_doc(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for name, size, color in [
        ("Heading 1", 16, RGBColor(46, 116, 181)),
        ("Heading 2", 13, RGBColor(46, 116, 181)),
        ("Heading 3", 12, RGBColor(31, 77, 120)),
    ]:
        st = styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = color
        st.paragraph_format.space_before = Pt(10)
        st.paragraph_format.space_after = Pt(5)


def add_title(doc, title, subtitle):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor(31, 77, 120)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(subtitle)
    r2.font.name = "Calibri"
    r2.font.size = Pt(11)
    r2.italic = True


def add_key_value_table(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    set_cell_text(hdr[0], "Field", True)
    set_cell_text(hdr[1], "Response", True)
    set_cell_shading(hdr[0], "F2F4F7")
    set_cell_shading(hdr[1], "F2F4F7")
    for k, v in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], k, True)
        set_cell_text(cells[1], v)
    doc.add_paragraph()


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Number")


def build_mahone():
    doc = Document()
    style_doc(doc)
    add_title(
        doc,
        "Proposal for Website Redesign and Support Services",
        "Town of Mahone Bay RFP2 | Submitted by Integrated IT Support Inc.",
    )
    add_key_value_table(
        doc,
        [
            ("Legal company name", COMPANY["legal_name"]),
            ("Primary contact", f"{COMPANY['contact']}, {COMPANY['title']}"),
            ("Address", COMPANY["address"]),
            ("Email / phone", f"{COMPANY['email']} | {COMPANY['phone']}"),
            ("Website", COMPANY["website"]),
            ("Proposal validity", "60 calendar days from RFP closing, subject to final contract terms."),
            ("Signature status", "Prepared for Ahmad Wasee signature before upload/submission."),
        ],
    )

    doc.add_heading("Cover Letter", level=1)
    for text in [
        "Integrated IT Support Inc. is pleased to submit this proposal to design, build, launch, train, and support a modern municipal website for the Town of Mahone Bay.",
        "IIS brings enterprise IT discipline from regulated financial services, municipal/public-sector support experience, Microsoft 365 and cloud administration, workflow automation, documentation, and practical AI-enabled support operations. For this project, IIS will keep the delivery simple for Town staff: clear planning, accessible design, clean content migration, staff training, and responsive post-launch support.",
        "Our proposed approach is intentionally practical. The Town receives an accessible, mobile-friendly, searchable, staff-editable website with integrations and support options that can grow over the three-year contract period without creating unnecessary technical lock-in.",
    ]:
        doc.add_paragraph(text)

    doc.add_heading("Understanding of the Project", level=1)
    add_bullets(
        doc,
        [
            "The current Town website was launched in 2020 and now needs a refreshed platform, structure, and content model.",
            "The new site must be a primary resource for residents, businesses, visitors, and internal staff.",
            "Town staff must be able to manage pages, documents, images, news, events, and updates without vendor dependency.",
            "The design must be responsive, easy to navigate, searchable, accessible, and owned by the Town.",
            "The solution must support or integrate with OneDrive, Brightly, Bids and Tenders, emergency alerts, business directory, analytics, document management, and service request functionality where feasible.",
        ],
    )

    doc.add_heading("Proposed Solution", level=1)
    doc.add_paragraph(
        "IIS proposes a managed municipal website rebuild using a mainstream content management platform configured for accessibility, maintainability, structured content, and staff usability. Final CMS selection will be confirmed during discovery based on Town preference, hosting constraints, content needs, security, and long-term supportability."
    )
    add_bullets(
        doc,
        [
            "Accessible, responsive public website with modern navigation, sitemap, search, and mobile-first layouts.",
            "Staff-editable CMS with roles/permissions, page templates, document library practices, media management, and publishing workflow.",
            "Content migration plan covering inventory, cleanup, page mapping, redirect planning, document handling, and Town review cycles.",
            "Integration readiness for third-party tools listed in the RFP, with lightweight API/embed/linking approach unless deeper integration is approved.",
            "Optional AI-assisted content governance and site-search enhancement for better public self-service without exposing private or sensitive data.",
        ],
    )

    doc.add_heading("Work Plan and Timeline", level=1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Phase", "Target Dates", "Key Activities", "Deliverables"]
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, True)
        set_cell_shading(table.rows[0].cells[i], "F2F4F7")
    rows = [
        ("1. Discovery", "Jul 6-Jul 12", "Kickoff, stakeholder review, current-site inventory, integration discovery.", "Project plan, content inventory, success criteria."),
        ("2. Architecture and design", "Jul 13-Jul 26", "Information architecture, wireframes, visual direction, accessibility plan.", "Sitemap, page templates, design approval."),
        ("3. Build and migration", "Jul 27-Aug 16", "CMS configuration, templates, navigation, search, forms/links, content migration.", "Working staging site, migrated priority content."),
        ("4. Review and training", "Aug 17-Aug 30", "Town review, accessibility checks, staff training, support manual.", "Training session, CMS manual, issue log."),
        ("5. Launch", "Sep 1-Sep 7", "Soft launch, QA, redirects, analytics, go-live support.", "Soft launch Sep 1, official launch Sep 7."),
        ("6. Support", "Post-launch", "Troubleshooting, updates, minor improvements, advisory support.", "Monthly support reporting and ticket log."),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, i == 0)
    doc.add_paragraph()

    doc.add_heading("Accessibility Compliance", level=1)
    add_bullets(
        doc,
        [
            "Design and build will target WCAG 2.1 AA-aligned practices and the Town's accessibility standards.",
            "Templates will use semantic headings, keyboard-friendly navigation, accessible color contrast, alt-text practices, readable layouts, and responsive behavior.",
            "IIS will provide an accessibility checklist and practical staff guidance so future content remains accessible after launch.",
        ],
    )

    doc.add_heading("Company Profile and Relevant Experience", level=1)
    add_bullets(
        doc,
        [
            "Integrated IT Support Inc. is an Ontario-based applied AI, automation, website, and managed IT services firm.",
            "Ahmad Wasee has 13+ years of enterprise IT engineering and support leadership across Raymond James, RBC Capital Markets, IBM, Ontario Health, Scotia Bank, City of Toronto, and legal/healthcare environments.",
            "Relevant capabilities include Microsoft 365, SharePoint/Teams/Outlook, Entra ID, Intune, ServiceNow ITSM, documentation, service desk leadership, workflow automation, security-aware support, training, and AI-enabled support systems.",
            "IIS operates ARIA, a production multi-agent AI operations platform used for procurement intelligence, knowledge retrieval, contract hunting, workflow automation, and operational briefings.",
            "For municipal web delivery, IIS will use a right-sized support bench for design, CMS build, accessibility review, and content migration while IIS retains project governance and accountability.",
        ],
    )

    doc.add_heading("Fee Structure", level=1)
    fee = doc.add_table(rows=1, cols=4)
    fee.style = "Table Grid"
    fee.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(["Item", "Description", "Price", "Notes"]):
        set_cell_text(fee.rows[0].cells[i], h, True)
        set_cell_shading(fee.rows[0].cells[i], "F2F4F7")
    fee_rows = [
        ("Base website redesign", "Discovery, design, CMS build, migration support, integrations setup, launch QA.", "CAD $68,500", "Firm fixed price, excluding HST."),
        ("Included training", "Two remote staff training sessions plus CMS/admin reference manual.", "Included", "Additional training at CAD $125/hour."),
        ("Standard support", "Email/phone troubleshooting, minor updates, CMS support, monthly health check.", "CAD $1,200/month", "Annual: CAD $14,400."),
        ("Enhanced support", "Standard support plus quarterly improvement sprint and analytics review.", "CAD $1,850/month", "Annual: CAD $22,200."),
        ("Optional AI content/search advisory", "Content governance, public-service FAQ mapping, AI-assisted search/readiness plan.", "CAD $4,800 fixed", "Optional add-on; no sensitive-data exposure."),
        ("Additional services", "Out-of-scope integrations, custom development, urgent work.", "CAD $125-$175/hour", "Quoted before work begins."),
    ]
    for row in fee_rows:
        cells = fee.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, i == 0)
    doc.add_paragraph("HST will be shown separately on invoice and is not included in the prices above.")

    doc.add_heading('Appendix "A" - Proposal Application Form', level=1)
    add_key_value_table(
        doc,
        [
            ("Company Legal Name", COMPANY["legal_name"]),
            ("Authorized Official", "Ahmad Wasee"),
            ("Date", date.today().isoformat()),
            ("Title", COMPANY["title"]),
            ("Print Name", "Ahmad Wasee"),
            ("Contact Person for Project", "Ahmad Wasee"),
            ("Address", COMPANY["address"]),
            ("Phone Number", COMPANY["phone"]),
            ("Fax Number", "N/A"),
            ("Email Address", COMPANY["email"]),
            ("Authorized Official Signature", "Pending Ahmad Wasee signature before final submission."),
        ],
    )

    path = OUT / "IIS_MahoneBay_RFP2_Website_Redesign_Proposal.docx"
    doc.save(path)
    return path


def build_ai_source_list():
    doc = Document()
    style_doc(doc)
    add_title(
        doc,
        "AI Source List Qualification Package",
        "CanadaBuys WS4286933967 | Integrated IT Support Inc.",
    )
    add_key_value_table(
        doc,
        [
            ("Opportunity", "Invitation to Qualify to Artificial Intelligence Source List"),
            ("Solicitation", "WS4286933967 / Doc4822970058"),
            ("Closing", "2026-09-30 14:00 EDT"),
            ("Submission portal", "SAP Business Network"),
            ("Supplier", COMPANY["legal_name"]),
            ("Contact", f"{COMPANY['contact']} | {COMPANY['email']} | {COMPANY['phone']}"),
        ],
    )
    doc.add_heading("Executive Positioning", level=1)
    doc.add_paragraph(
        "Integrated IT Support Inc. provides practical AI, automation, managed IT, and Microsoft-cloud enablement services for organizations that need AI to improve real operational workflows without creating unmanaged risk. IIS is best positioned for responsible AI advisory, AI readiness, workflow automation, knowledge-base and document automation, M365/Copilot readiness, helpdesk automation, and AI-enabled support operations."
    )
    doc.add_heading("Relevant Qualifications", level=1)
    add_bullets(
        doc,
        [
            "Production ARIA platform: scheduled autonomous agents for procurement intelligence, knowledge retrieval, operational briefings, lead capture, and finance/tax workflow support.",
            "Hands-on LLM stack: Anthropic Claude, OpenAI GPT-4/GPT-4o, Google Gemini, Azure OpenAI, Microsoft Copilot/Copilot Studio, MCP, tool/function calling, RAG patterns, structured outputs, and prompt regression practices.",
            "Enterprise IT background: 13+ years across regulated financial services, healthcare, public-sector support, service desk leadership, major incident management, change/problem management, CMDB, and ITSM.",
            "Microsoft environment strength: M365, SharePoint, Teams, Outlook, Entra ID, Intune, Defender, Azure Monitor, Azure/OpenAI, Conditional Access, MFA, hybrid identity, and endpoint lifecycle practices.",
            "Governance-aware delivery: experience in OSFI/IIROC, SOX, PIPEDA, PHIPA, audit-ready change management, access controls, documentation, and evidence-of-control work.",
        ],
    )
    doc.add_heading("Service Categories to Claim Where Allowed", level=1)
    rows = [
        ("AI readiness and strategy", "AI use-case discovery, risk triage, operating model, roadmap, stakeholder workshops."),
        ("Responsible AI governance", "Policy templates, model/use-case registers, human review practices, prompt/data handling controls."),
        ("Workflow and document automation", "Intake automation, document drafting, knowledge retrieval, structured outputs, approval workflows."),
        ("AI service desk enablement", "Ticket triage, runbook drafting, incident summaries, knowledge-base optimization, analyst coaching."),
        ("M365/Copilot readiness", "Readiness assessment, permissions review, adoption training, SharePoint/Teams knowledge structure."),
        ("Managed AI operations", "Monitoring, prompt regression/evals, change control, documentation, operational reporting."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    for i, h in enumerate(["Category", "IIS Response Position"]):
        set_cell_text(table.rows[0].cells[i], h, True)
        set_cell_shading(table.rows[0].cells[i], "F2F4F7")
    for row in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], row[0], True)
        set_cell_text(cells[1], row[1])
    doc.add_paragraph()
    doc.add_heading("No-Fabrication Qualification Matrix", level=1)
    matrix = [
        ("Canadian supplier identity", "Known", "Integrated IT Support Inc.; address/contact confirmed from Sourcewell package."),
        ("Responsible AI experience", "Supported", "ARIA, MCP, LLM workflows, governance-aware delivery; avoid claiming formal AI audits unless documented."),
        ("Security certifications", "Unknown", "Do not claim SOC 2, ISO 27001, or similar unless certificate is provided."),
        ("Insurance", "Unknown", "Confirm CGL/professional/cyber coverage before any legal attestation."),
        ("Public-sector AI references", "Limited", "Use regulated enterprise/public-sector IT examples carefully; do not list contacts without permission."),
        ("Subcontractor capacity", "Allowed if identified", "Use support-bench language only where permitted and name partners once confirmed."),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for i, h in enumerate(["Criterion", "Status", "Action"]):
        set_cell_text(table.rows[0].cells[i], h, True)
        set_cell_shading(table.rows[0].cells[i], "F2F4F7")
    for row in matrix:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, i == 0)
    doc.add_paragraph()
    doc.add_heading("Portal Copy Block", level=1)
    doc.add_paragraph(
        "Integrated IT Support Inc. is an Ontario-based applied AI and managed IT services firm led by Ahmad Wasee, a senior AI engineer and enterprise IT professional with 13+ years of experience in regulated financial services, healthcare, public-sector support, ITSM, Microsoft 365, Azure, and automation. IIS designs and operates ARIA, a production multi-agent AI operations platform using LLMs, MCP tool orchestration, workflow automation, structured knowledge retrieval, and governance-aware delivery practices. IIS supports responsible AI adoption through readiness assessments, use-case discovery, governance templates, M365/Copilot readiness, document and workflow automation, helpdesk automation, analyst enablement, training, and managed AI operations."
    )
    doc.add_heading("Before Final SAP Submission", level=1)
    add_numbered(
        doc,
        [
            "Log in to SAP Business Network from the CanadaBuys notice.",
            "Download the current qualification form and mandatory criteria from the live event.",
            "Map each mandatory criterion to this package and the Sourcewell material.",
            "Confirm insurance, legal attestations, security certifications, and reference permissions.",
            "Upload only final documents that Ahmad has reviewed and approved.",
        ],
    )
    path = OUT / "IIS_CanadaBuys_AI_Source_List_Qualification_Package.docx"
    doc.save(path)
    return path


def build_uncommon():
    doc = Document()
    style_doc(doc)
    add_title(
        doc,
        "Uncommon Schools Managed IT Services Pursuit Package",
        "RFP 2026-06-USI | Integrated IT Support Inc.",
    )
    add_key_value_table(
        doc,
        [
            ("RFP", "Uncommon Schools Managed IT Services, RFP 2026-06-USI"),
            ("Due", "2026-06-23 at 3:00 PM Eastern Time"),
            ("Submission", "Email PDF proposal plus separate cost workbook/signed PDF to msp.rfp.2026@uncommonschools.org"),
            ("Recommended posture", "Prime only with named state-compliant onsite/MSP partners; otherwise subcontract/workstream partner."),
            ("Current status", "Do not final-submit until partner, insurance, certifications, state licensing, and references are confirmed."),
        ],
    )
    doc.add_heading("Why This Is Attractive", level=1)
    add_bullets(
        doc,
        [
            "Large, multi-region managed IT opportunity covering enterprise infrastructure, cloud/device management, cybersecurity, end-user support, onsite support, classroom technology, lifecycle management, BCDR, project management, and vendor advocacy.",
            "IIS has strong alignment in service desk leadership, ITSM, Microsoft 365, Entra/Intune, incident management, automation, training, documentation, and AI-enabled support operations.",
            "ARIA and Ahmad's regulated enterprise experience can differentiate IIS on automation, reporting, escalation quality, operational documentation, and AI-assisted support improvement.",
        ],
    )
    doc.add_heading("Hard Submission Gates", level=1)
    add_bullets(
        doc,
        [
            "State licensing for New York, New Jersey, and Massachusetts must be in place by contract start, and proof must be available before contract execution.",
            "Documentation of engineers holding role-relevant certifications must be submitted with the proposal.",
            "Cybersecurity section asks for SOC 2 Type II or ISO 27001 proof, or sub-service provider certifications where those services are subcontracted.",
            "The proposal requires current insurance representations, including cyber errors and omissions coverage.",
            "Forms require authorized signature, felony/non-collusion/ownership/political contribution/debarment certifications, and New Jersey business/authority documentation where applicable.",
            "References require company names, contacts, phone numbers, and service descriptions; do not list employers/clients as references without permission.",
        ],
    )
    doc.add_heading("IIS Best-Fit Workstreams", level=1)
    rows = [
        ("Service desk quality and escalation", "L1-L3 practices, SLA/KPI dashboards, major incident discipline, analyst coaching."),
        ("M365 / Entra / Intune", "Identity, endpoint, conditional access, Microsoft 365 administration and support."),
        ("AI-enabled support automation", "Knowledge retrieval, triage templates, incident summaries, runbook drafting, reporting."),
        ("ITSM and documentation", "ServiceNow-style workflows, CMDB discipline, knowledge base, onboarding/offboarding packages."),
        ("Security-aware operations", "Access controls, Defender exposure, vulnerability coordination, audit-ready evidence practices."),
        ("Transition support", "Knowledge transfer, runbook development, asset/process inventory, steady-state handoff."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    for i, h in enumerate(["Workstream", "IIS Position"]):
        set_cell_text(table.rows[0].cells[i], h, True)
        set_cell_shading(table.rows[0].cells[i], "F2F4F7")
    for row in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], row[0], True)
        set_cell_text(cells[1], row[1])
    doc.add_paragraph()
    doc.add_heading("Partner Support Statement", level=1)
    doc.add_paragraph(
        "IIS will manage the AI-enabled support operations, service desk quality, M365/identity/endpoint workflow, documentation, reporting, and transition-improvement components. For onsite field coverage, state licensing, education MSP capacity, SOC/MDR services, and jurisdiction-specific school compliance, IIS will use named qualified partners and subcontractors under IIS delivery governance. Partners must provide current certifications, insurance evidence, role coverage, and state-compliance documentation before final proposal submission."
    )
    doc.add_heading("No-Send Partner Pitch", level=1)
    doc.add_paragraph(
        "Integrated IT Support Inc. is preparing a response/workstream support package for Uncommon Schools RFP 2026-06-USI. IIS brings enterprise ITSM, M365/Entra/Intune, service desk leadership, AI-enabled support automation, documentation, reporting, and regulated-environment experience. We are seeking a qualified education MSP/onsite support partner with NY/NJ/MA coverage, state licensing path, SOC/ISO or MDR partner certifications, education references, and field technician capacity. IIS can strengthen the proposal through automation, transition documentation, escalation quality, and AI-enabled service delivery while the partner anchors onsite and school-specific compliance."
    )
    doc.add_heading("Decision", level=1)
    doc.add_paragraph(
        "Proceed as a high-value pursuit, but do not send a direct prime proposal until IIS has named compliant delivery partners and Ahmad confirms the legal certifications, insurance position, references, and signature-ready forms. A weak direct submission would be less competitive than a strong partner-backed submission."
    )
    path = OUT / "IIS_Uncommon_RFP2026_06_USI_Pursuit_and_Partner_Package.docx"
    doc.save(path)
    return path


def build_submission_readme(paths):
    readme = OUT / "SUBMISSION_NEXT_STEPS.md"
    readme.write_text(
        "\n".join(
            [
                "# Selected Pursuits - Submission Next Steps",
                "",
                "Created: 2026-06-05",
                "",
                "## 1. Mahone Bay RFP2",
                "",
                f"- Proposal file: `{paths['mahone']}`",
                "- Submit via Bids&Tenders after login/registration for the bid.",
                "- Final gate: Ahmad must review pricing, confirm HST/tax treatment, and sign Appendix A.",
                "- Portal: https://townofmahonebay.bidsandtenders.ca/Module/Tenders/en/Tender/Detail/d34bda17-ea5d-407f-997f-5b269a9b1e78",
                "",
                "## 4. CanadaBuys AI Source List",
                "",
                f"- Qualification package: `{paths['ai']}`",
                "- Submit through SAP Business Network from the CanadaBuys notice.",
                "- Final gate: download live SAP event forms/mandatory criteria, confirm insurance/security/legal attestations, and upload only reviewed final documents.",
                "- Portal: https://canadabuys.canada.ca/en/tender-opportunities/tender-notice/ws4286933967-doc4822970058",
                "",
                "## 6. Uncommon Schools Managed IT",
                "",
                f"- Pursuit/partner package: `{paths['uncommon']}`",
                "- Downloaded RFP docs are in `procurement-downloads/uncommon-2026-06-usi`.",
                "- Final gate: do not submit as prime until state licensing path, SOC/ISO/MDR certifications, insurance, engineer certification counts, references, and NJ forms are ready.",
                "- Submission email in RFP: msp.rfp.2026@uncommonschools.org",
                "- Portal/page: https://uncommonschools.org/current-legal-notices/",
                "",
            ]
        ),
        encoding="utf-8",
    )
    return readme


def main():
    paths = {
        "mahone": build_mahone(),
        "ai": build_ai_source_list(),
        "uncommon": build_uncommon(),
    }
    paths["readme"] = build_submission_readme(paths)
    for k, v in paths.items():
        print(f"{k}: {v}")


if __name__ == "__main__":
    main()
