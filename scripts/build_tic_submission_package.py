from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parents[1]
SRC = (
    ROOT
    / "procurement-downloads"
    / "tic-it-advisory-2025-42r-browser"
    / "TIC_TEC_MUL_RFQ_2025_42_R_IT_Advisors_Appendix_B_Resource_Qualification_Response_Form.docx"
)
OUT = ROOT / "artifacts" / "tic-it-advisory-submission"
APPENDIX_B = OUT / "IIS_TIC_TEC_MUL_RFQ_2025_42R_Appendix_B_Ahmad_Wasee_Technical_Analyst.docx"
RESUME_PDF = OUT / "IIS_TIC_Ahmad_Wasee_Technical_Analyst_Resume_4pg.pdf"
RATE_PDF = OUT / "IIS_TIC_Ahmad_Wasee_Technical_Analyst_Hourly_Rate.pdf"


def clear_cell(cell):
    for p in cell.paragraphs:
        p.text = ""


def set_cell(cell, text: str):
    clear_cell(cell)
    p = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(9)


def append_para(doc: Document, text: str, bold: bool = False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(10)
    return p


def build_appendix_b() -> Path:
    OUT.mkdir(parents=True, exist_ok=True)
    doc = Document(str(SRC))

    # Respondent Identification
    t = doc.tables[2]
    values = {
        1: "Integrated IT Support Inc.",
        2: "30 Fothergill Court, Whitby, Ontario L1P 1L4, Canada",
        3: "+1 647-581-3182",
        4: "https://iisupp.net",
        5: "Ahmad Wasee",
        6: "Founder and Senior AI Engineer",
        7: "ahmad.wasee@iisupp.net",
        8: "+1 647-581-3182",
        11: "Ahmad Wasee",
        12: "Founder and Senior AI Engineer",
        13: "ahmad.wasee@iisupp.net",
        14: "+1 647-581-3182",
        15: "30 Fothergill Court, Whitby, Ontario L1P 1L4, Canada",
    }
    for row, value in values.items():
        set_cell(t.rows[row].cells[-1], value)

    # Resource Identification
    set_cell(doc.tables[3].rows[2].cells[-1], "Ahmad Wasee - proposed Service Area: Technical Analyst")

    # Certifications / education
    certs = [
        ["Advanced Diploma - Computer Networking & Technical Support", "Seneca College of Applied Arts & Technology", "2011"],
        ["ITIL Foundations", "RBC", "2017"],
        ["Six Sigma Yellow Belt", "RBC", "2017"],
        ["Project Management Certification", "Training provider / continuing education", "2023"],
        ["Agile Project Management Certification", "Training provider / continuing education", "2023"],
    ]
    t = doc.tables[6]
    for row_values in certs:
        cells = t.add_row().cells
        set_cell(cells[0], row_values[0])
        set_cell(cells[1], row_values[1])
        set_cell(cells[3], row_values[2])

    # Demonstrated experience
    projects = [
        {
            "name": "Raymond James - Regulated Enterprise ITSM, Endpoint, Executive/Trader-Floor Support, and Automation",
            "dates": "Feb 2020 - Present",
            "description": (
                "Resource role: Senior Technical Analyst - AI & Automation Engineering. Ahmad supports a regulated "
                "capital-markets IT environment with ServiceNow ITSM workflow, Microsoft 365, Entra ID, Intune, Active "
                "Directory, Citrix, Windows 11/Server, executive and trader-floor support, major-incident discipline, "
                "runbook/knowledge-base authoring, CMDB practices, endpoint migration, incident-summary/RCA drafting, "
                "and support automation. Relevance to Technical Analyst service area: 3+ years hands-on enterprise IT "
                "support, troubleshooting, user/application support, documentation, escalation, endpoint/M365 support, "
                "and operational process improvement in a high-availability regulated environment."
            ),
        },
        {
            "name": "RBC Capital Markets - Global Service Desk Leadership, Major Incident, Reporting, and Automation",
            "dates": "Mar 2013 - Dec 2017",
            "description": (
                "Resource role: CM Senior Technical Support Analyst / Team Lead / Acting Manager. Ahmad supported "
                "Capital Markets, Wealth Management, Investor and Treasury Services clients with hardware, software, "
                "application and access issues in a 24x7 global environment. He managed ServiceNow/HP Service Manager "
                "queues, escalations, SLA/KPI reporting, CMDB lifecycle practices, major incidents, on-call support, "
                "analyst coaching, and automation. Relevance to Technical Analyst service area: senior support and "
                "technical analyst experience across complex enterprise systems, incident management, reporting, "
                "automation, documentation, and stakeholder communication."
            ),
        },
        {
            "name": "Integrated IT Support Inc. - ARIA AI Automation, Procurement Intelligence, and Support Workflow Platform",
            "dates": "2023 - Present",
            "description": (
                "Resource role: Founder and Senior AI Engineer. Ahmad designed and operates ARIA, a production "
                "multi-agent automation platform using Claude/OpenAI tooling, MCP orchestration, Python, Node.js, "
                "browser automation, Netlify, markdown/YAML knowledge architecture, scheduled agents, and CRM/workflow "
                "automation. Relevance to Technical Analyst service area: hands-on IT automation, troubleshooting, "
                "knowledge management, support-process tooling, technical documentation, web/API integration, and "
                "practical AI-assisted service support."
            ),
        },
    ]
    t = doc.tables[8]
    set_cell(t.rows[1].cells[0], projects[0]["name"])
    set_cell(t.rows[2].cells[0], projects[0]["dates"].split(" - ")[0])
    set_cell(t.rows[2].cells[1], projects[0]["dates"].split(" - ")[1])
    set_cell(t.rows[4].cells[0], projects[0]["description"])
    set_cell(t.rows[6].cells[0], projects[1]["name"])
    set_cell(t.rows[7].cells[0], projects[1]["dates"].split(" - ")[0])
    set_cell(t.rows[7].cells[1], projects[1]["dates"].split(" - ")[1])
    set_cell(t.rows[9].cells[0], projects[1]["description"])

    t = doc.tables[9]
    set_cell(t.rows[1].cells[0], projects[2]["name"])
    set_cell(t.rows[2].cells[0], projects[2]["dates"].split(" - ")[0])
    set_cell(t.rows[2].cells[1], projects[2]["dates"].split(" - ")[1])
    set_cell(t.rows[4].cells[0], projects[2]["description"])

    # Unique expertise
    unique = (
        "Ahmad combines senior enterprise technical support experience with modern AI/workflow automation. "
        "For TIC's Technical Analyst service area, his distinctive value is the ability to resolve support issues, "
        "document repeatable fixes, improve ITSM workflows, create operational dashboards/reports, and translate "
        "technical patterns into runbooks, user support notes, and automation candidates. He has worked in regulated "
        "financial services environments where availability, evidence of control, executive communication, and clear "
        "escalation practices matter. IIS will not represent Ahmad as an OnBase/Hyland certified resource or BC MOTI "
        "specialist; this response is intentionally scoped to Technical Analyst services."
    )
    set_cell(doc.tables[10].add_row().cells[0], unique)

    set_cell(
        doc.tables[11].add_row().cells[0],
        "N/A for proposed Service Area: Technical Analyst. Work samples are requested for Corporate Technology Trainer, Senior IT Project Manager, Senior Business Analyst, Business Analyst, and Enterprise/Technical/Solutions Architect service areas. Ahmad is proposed only as a Technical Analyst. Resume and demonstrated experience are provided for verification.",
    )
    set_cell(
        doc.tables[12].add_row().cells[0],
        "Resume submitted as separate file: IIS_TIC_Ahmad_Wasee_Technical_Analyst_Resume_4pg.pdf",
    )
    set_cell(
        doc.tables[13].add_row().cells[0],
        "Hourly rate submitted as separate file: IIS_TIC_Ahmad_Wasee_Technical_Analyst_Hourly_Rate.pdf. Proposed all-inclusive rate: CAD 125.00/hour for Ahmad Wasee, Technical Analyst.",
    )

    # Complete the email submission declaration at the end of the official form.
    append_para(doc, "")
    p = append_para(doc, "Submission Declaration - Completed for Email Submission", True)
    p.runs[0].font.size = Pt(12)
    append_para(doc, "Signature of authorized representative: Ahmad Wasee")
    append_para(doc, "Print your name: Ahmad Wasee")
    append_para(doc, "Print name of your authorized representative: Ahmad Wasee")
    append_para(doc, f"Date: {date.today().strftime('%B %d, %Y')}")
    append_para(
        doc,
        "By submitting this response by email, Integrated IT Support Inc. agrees to the RFQ terms and confirms that Ahmad Wasee's typed name constitutes the legally binding signature of the authorized representative.",
    )

    doc.save(APPENDIX_B)
    return APPENDIX_B


def styles():
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("Title", parent=base["Title"], fontName="Helvetica-Bold", fontSize=18, leading=22, textColor=colors.HexColor("#163A5F"), alignment=0),
        "h": ParagraphStyle("H", parent=base["Heading2"], fontName="Helvetica-Bold", fontSize=12, leading=15, textColor=colors.HexColor("#1F4D78"), spaceBefore=8, spaceAfter=4),
        "body": ParagraphStyle("Body", parent=base["BodyText"], fontName="Helvetica", fontSize=9.5, leading=12, spaceAfter=4),
        "small": ParagraphStyle("Small", parent=base["BodyText"], fontName="Helvetica", fontSize=8.5, leading=10.5, spaceAfter=3),
    }


def para(text, style):
    return Paragraph(text, style)


def bullets(items, style):
    return ListFlowable([ListItem(para(i, style), leftIndent=10) for i in items], bulletType="bullet", leftIndent=14, bulletIndent=4)


def build_resume() -> Path:
    s = styles()
    doc = SimpleDocTemplate(str(RESUME_PDF), pagesize=LETTER, leftMargin=0.65 * inch, rightMargin=0.65 * inch, topMargin=0.6 * inch, bottomMargin=0.6 * inch)
    story = [
        para("Ahmad Wasee", s["title"]),
        para("Technical Analyst | Enterprise IT Support, ITSM, Microsoft 365/Endpoint, Automation", s["body"]),
        para("Integrated IT Support Inc. | Whitby, Ontario | ahmad.wasee@iisupp.net | +1 647-581-3182 | https://iisupp.net", s["small"]),
        Spacer(1, 6),
        para("Profile", s["h"]),
        para("Senior enterprise technical analyst and founder of Integrated IT Support Inc. with 13+ years supporting regulated financial services, public-sector and healthcare-adjacent environments. Strong fit for TIC's Technical Analyst service area: desktop/user/application support, Microsoft 365, endpoint/identity support, ServiceNow/ITSM workflow, incident/change/problem practices, executive support, documentation, reporting, and automation.", s["body"]),
        para("Relevant Skills", s["h"]),
        bullets([
            "Enterprise support: L1/L2/L3 troubleshooting, executive/trader-floor support, incident triage, escalation, major incident management, SLA/KPI reporting.",
            "Microsoft/endpoint: Microsoft 365, Teams, Outlook, SharePoint, Entra ID, Intune, Active Directory, Group Policy, Windows 10/11, Windows Server, Citrix.",
            "ITSM/process: ServiceNow ITSM/ITOM, HP Service Manager, CMDB practices, CSDM alignment, change/problem management, runbooks and knowledge-base authoring.",
            "Automation/documentation: PowerShell, Python, VBA, VB.NET, Excel automation, reporting dashboards, AI-assisted runbooks, process improvement.",
            "Security/privacy-aware support: SOX, OSFI, IIROC, PIPEDA, PHIPA, RSA Archer, Microsoft Defender, access controls and audit-ready change evidence.",
        ], s["body"]),
        para("Professional Experience", s["h"]),
        para("<b>Founder and Senior AI Engineer, Integrated IT Support Inc.</b> | 2023-present", s["body"]),
        bullets([
            "Designed and operate ARIA, a production AI automation and procurement-intelligence platform using Claude/OpenAI tooling, MCP orchestration, Python, Node.js, Netlify, browser automation and file-based knowledge architecture.",
            "Built repeatable support, procurement, prospecting, knowledge retrieval, reporting and tax/expense automation workflows.",
            "Created runbooks, prompt/task pipelines, support knowledge artifacts and operating procedures for AI-assisted business operations.",
        ], s["body"]),
        para("<b>Senior Technical Analyst - AI & Automation Engineering, Raymond James</b> | Feb 2020-present", s["body"]),
        bullets([
            "Support regulated capital-markets IT operations including executive/trader-floor support, incident response, endpoint migration, ServiceNow ITSM workflow, CMDB practices and support documentation.",
            "Hands-on Microsoft 365, Entra ID, Intune, Active Directory, Citrix, Windows 11/Server, Azure and knowledge-base/runbook development.",
            "Applied AI and automation practices for incident summaries, RCA drafting, support knowledge structure, ticket templates and recurring support-pattern reduction.",
        ], s["body"]),
        para("<b>Executive Support - Senior Technical Analyst, Scotia Bank</b> | Nov 2019-Feb 2020", s["body"]),
        bullets([
            "Supported Windows 10 and Office 365 migration issues, executive support and high-volume laptop deployment with quality and speed.",
        ], s["body"]),
        para("<b>IT Consultant / Specialist, Ontario Health</b> | Jul 2019-Nov 2019", s["body"]),
        bullets([
            "Supported digital health solutions, system configurations, software installations, performance optimization, troubleshooting, QA/UAT, documentation and PHIPA-aware support practices.",
        ], s["body"]),
        para("<b>Tier 1/2/3 Support Analyst, Cavalluzzo LLP</b> | Nov 2018-Jul 2019", s["body"]),
        bullets([
            "Supported lawyers and staff with hardware, software, access, mobile, networking, domain/group policy, server/switch monitoring, ServiceNow tasks and document-system migrations.",
        ], s["body"]),
        para("<b>CM Senior Technical Support Analyst / Team Lead / Acting Manager, RBC Capital Markets</b> | Mar 2013-Dec 2017", s["body"]),
        bullets([
            "Led 24x7 global support queues for capital markets and wealth clients; managed escalations, major incidents, SLA/KPI reporting, CMDB practices, analyst coaching and on-call support.",
            "Built automation including stale-ticket closure and diagnostic tooling, improving support efficiency and consistency.",
        ], s["body"]),
        para("<b>Senior Technical Support Specialist, City of Toronto</b> | Nov 2012-Jan 2013", s["body"]),
        bullets([
            "Supported enterprise infrastructure monitoring, directories, messaging, system performance and team coordination in a municipal environment.",
        ], s["body"]),
        para("<b>Senior Technical Analyst / Incident Queue Specialist, IBM</b> | Jan 2011-Nov 2012", s["body"]),
        bullets([
            "Managed incident queues and resolved backup/restore, Microsoft Office, WebEx, Windows application, account provisioning and mobile device support issues.",
        ], s["body"]),
        para("Education and Certifications", s["h"]),
        bullets([
            "Advanced Diploma, Computer Networking & Technical Support, Seneca College of Applied Arts & Technology.",
            "ITIL Foundations; Six Sigma Yellow Belt; Project Management Certification; Agile Project Management Certification.",
            "Anthropic Academy AI/Claude coursework completed/in progress in 2026.",
        ], s["body"]),
    ]
    doc.build(story)
    return RESUME_PDF


def build_rate() -> Path:
    s = styles()
    doc = SimpleDocTemplate(str(RATE_PDF), pagesize=LETTER, leftMargin=0.8 * inch, rightMargin=0.8 * inch, topMargin=0.8 * inch, bottomMargin=0.8 * inch)
    data = [
        ["Opportunity", "TIC TEC MUL RFQ 2025 42 R - IT Advisory Services"],
        ["Respondent", "Integrated IT Support Inc."],
        ["Resource", "Ahmad Wasee"],
        ["Service Area", "Technical Analyst"],
        ["Hourly Rate", "CAD 125.00/hour"],
        ["Rate Basis", "All-inclusive hourly rate, including fees, business/admin overhead, profit, incidentals, routine out-of-pocket costs and disbursements. No separate travel, administration or disbursement charges will be added unless expressly authorized under the RFQ or resulting contract terms."],
    ]
    table = Table([[para(a, s["small"]), para(b, s["small"])] for a, b in data], colWidths=[1.7 * inch, 4.8 * inch])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F2F4F7")),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CCCCCC")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story = [para("Hourly Rate", s["title"]), para("Separate rate file for TIC TEC MUL RFQ 2025 42 R", s["body"]), Spacer(1, 12), table]
    doc.build(story)
    return RATE_PDF


def main():
    print(build_appendix_b())
    print(build_resume())
    print(build_rate())


if __name__ == "__main__":
    main()
